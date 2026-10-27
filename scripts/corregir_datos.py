"""Genera "Clases Particulares - corregido.xlsx" a partir de la copia original del Sheet.

Uso (desde la raíz del repo):
    .venv\\Scripts\\python scripts\\corregir_datos.py
    .venv\\Scripts\\python scripts\\corregir_datos.py --entrada X.xlsx --salida Y.xlsx --corte 2026-10-26

El original NO se modifica. Las decisiones sobre registros dudosos están en CORRECCIONES,
REVISAR y CONFIRMADOS (respuestas de Federico del 26/10/2026). Cada cambio queda en la hoja LEEME.
"""
import argparse
import sys
from datetime import date, timedelta
from pathlib import Path

import openpyxl
import pandas as pd

RAIZ = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(RAIZ))

from logica import (  # noqa: E402
    COLUMNAS_ALUMNOS, COLUMNAS_CLASES, COLUMNAS_CLASES_VIEJAS, COLUMNAS_IPC, COLUMNAS_TARIFAS,
    normalizar_texto, parse_fecha, parse_id, parse_numero, texto,
)

# ─────────────────────────────────────────────
# DECISIONES (sección 3 del pedido)
# ─────────────────────────────────────────────
# ID_Clase → (MontoCobrado real, motivo). Total se iguala a MontoCobrado.
CORRECCIONES = {
    7: (14000, "Redondeo: 3 h × $4.667 daba $14.001; monto real $14.000 (confirmado)"),
    22: (10000, "Redondeo: 1,5 h × $6.666 daba $9.999; monto real $10.000 (confirmado)"),
    26: (10000, "Redondeo: 1,5 h × $6.666 daba $9.999; monto real $10.000 (confirmado)"),
    25: (13000, "2 h × $6.666 daba $13.332; monto real $13.000 (confirmado)"),
}
# ID_Clase → nota. Quedan con los datos cargados y Revisar = Sí.
REVISAR = {
    43: "Cargada 1 h × $10.000. Sospecha: 2 h × $5.000. Sin confirmar: revisar horas reales.",
    63: "Cargada 2,5 h × $4.000 = $10.000. Sospecha: 2 h × $5.000. Sin confirmar: revisar horas reales.",
    64: "Cargada 2,5 h × $4.000 = $10.000 con ficha de $6.000 ($4.000 era el precio por defecto del "
        "formulario). Sin confirmar: revisar monto u horas.",
}
# ID_Clase → motivo. Se dejan como están (se registra en LEEME para la auditoría).
CONFIRMADOS = {
    56: "3 h, Total $20.000: el monto es real ($6.666/h era un precio forzado). Sin cambios.",
    58: "2 h × $4.000 con Total $20.000: confirmado que se cobró $20.000. Sin cambios.",
    66: "Misma fecha y alumno que la 67: confirmado que son dos clases distintas. Sin cambios.",
    67: "Misma fecha y alumno que la 66: confirmado que son dos clases distintas. Sin cambios.",
}
# Alumnos con alta reciente que quedan Activos aunque no tengan clases en la ventana
ACTIVOS_FORZADOS = {14: "Alta reciente (17/10/2026)"}
DIAS_INACTIVIDAD = 60

COLUMNAS_TEXTO_ALUMNOS = {'Nombre', 'Materia', 'Curso', 'Teléfono', 'Contacto', 'Día', 'Frecuencia',
                          'Modalidad', 'FechaAlta', 'Estado', 'Observaciones'}


def leer_hoja(wb, nombre):
    ws = wb[nombre]
    filas = list(ws.iter_rows(values_only=True))
    enc = [normalizar_texto(h) for h in filas[0] if h is not None]
    datos = [{normalizar_texto(h): v for h, v in zip(filas[0], f) if h is not None}
             for f in filas[1:] if any(v not in (None, '') for v in f)]
    return enc, datos


def fecha_iso(v):
    f = parse_fecha(v)
    return f.date().isoformat() if not pd.isna(f) else ''


def num(v, tipo='dinero'):
    n = parse_numero(v, tipo=tipo)
    if n is None:
        return None
    return int(n) if float(n).is_integer() else round(n, 2)


def fmt(v):
    return '' if v is None else str(v)


def corregir(entrada, salida, corte):
    wb = openpyxl.load_workbook(entrada, data_only=True)
    enc_clases, clases_raw = leer_hoja(wb, 'Clases')
    enc_alumnos, alumnos_raw = leer_hoja(wb, 'Alumnos')
    faltan = [c for c in COLUMNAS_CLASES_VIEJAS if c not in enc_clases]
    assert not faltan, f"Faltan columnas en Clases: {faltan}"
    leeme = []

    def log(hoja, id_, columna, antes, despues, motivo):
        leeme.append([hoja, id_, columna, fmt(antes), fmt(despues), motivo])

    log('Todas', '', '(todas)', 'números como texto ("4000.0")', 'números reales',
        'Conversión de tipos: IDs, horas y montos pasan a ser números; fechas quedan como texto ISO.')
    log('Clases', '', 'MontoCobrado, HorasReales, Revisar, NotaRevision', '(no existían)', '(agregadas al final)',
        'Regla general: MontoCobrado = Total y HorasReales = Duración(h), salvo los registros listados abajo.')

    # ── Clases ──
    clases = []
    for r in clases_raw:
        c = {
            'ID_Clase': parse_id(r['ID_Clase']),
            'Fecha': fecha_iso(r['Fecha']),
            'ID_Alumno': parse_id(r['ID_Alumno']),
            'Alumno': texto(r['Alumno']),
            'Materia': texto(r['Materia']),
            'Duración(h)': num(r['Duración(h)'], 'horas'),
            'PrecioHora': num(r['PrecioHora']),
            '¿Pagada?': texto(r['¿Pagada?']),
            'Método': texto(r['Método']),
            'Total': num(r['Total']),
        }
        c['MontoCobrado'] = num(r.get('MontoCobrado')) if r.get('MontoCobrado') not in (None, '') else c['Total']
        c['HorasReales'] = (num(r.get('HorasReales'), 'horas') if r.get('HorasReales') not in (None, '')
                            else c['Duración(h)'])
        c['Revisar'] = texto(r.get('Revisar'))
        c['NotaRevision'] = texto(r.get('NotaRevision'))
        id_ = c['ID_Clase']
        if id_ in CORRECCIONES:
            monto, motivo = CORRECCIONES[id_]
            log('Clases', id_, 'MontoCobrado', c['MontoCobrado'], monto, motivo)
            if c['Total'] != monto:
                log('Clases', id_, 'Total', c['Total'], monto, 'Total se iguala a MontoCobrado (compatibilidad)')
            c['MontoCobrado'] = c['Total'] = monto
        if id_ in REVISAR:
            log('Clases', id_, 'Revisar / NotaRevision', '', 'Sí', REVISAR[id_])
            c['Revisar'], c['NotaRevision'] = 'Sí', REVISAR[id_]
        if id_ in CONFIRMADOS:
            log('Clases', id_, '(ninguna)', '', '', CONFIRMADOS[id_])
        clases.append(c)

    ids_vistos = set()
    for c in clases:
        assert c['ID_Clase'] not in ids_vistos, f"ID_Clase repetido: {c['ID_Clase']}"
        ids_vistos.add(c['ID_Clase'])
        assert c['HorasReales'] and c['HorasReales'] > 0, f"Horas inválidas en clase {c['ID_Clase']}"
    for id_ in list(CORRECCIONES) + list(REVISAR) + list(CONFIRMADOS):
        assert id_ in ids_vistos, f"El ID_Clase {id_} de las decisiones no existe en el Excel"

    # ── Tarifas inferidas del historial ──
    alumnos = []
    for r in alumnos_raw:
        a = {col: texto(r.get(col)) for col in COLUMNAS_ALUMNOS}
        a['ID_Alumno'] = parse_id(r['ID_Alumno'])
        a['PrecioHora'] = num(r['PrecioHora']) or 0
        a['FechaAlta'] = fecha_iso(r.get('FechaAlta'))
        alumnos.append(a)

    df = pd.DataFrame(clases)
    df['_fecha'] = pd.to_datetime(df['Fecha'])
    tarifas = []
    for a in alumnos:
        hist = df[(df['ID_Alumno'] == a['ID_Alumno']) & (df['Revisar'] != 'Sí')].sort_values(['_fecha', 'ID_Clase'])
        previa = None
        tasas = []
        for _, c in hist.iterrows():
            tasa = round(c['MontoCobrado'] / c['HorasReales'])
            if tasa != previa:
                tarifas.append({'ID_Alumno': a['ID_Alumno'], 'Desde': c['Fecha'], 'Tipo': 'Hora', 'Monto': tasa,
                                'Nota': f"inferida (clase {c['ID_Clase']})"})
                previa = tasa
            tasas.append(tasa)
        ficha = round(a['PrecioHora'])
        if previa is None:
            tarifas.append({'ID_Alumno': a['ID_Alumno'], 'Desde': a['FechaAlta'] or corte.isoformat(),
                            'Tipo': 'Hora', 'Monto': ficha, 'Nota': 'inferida (ficha, sin clases)'})
        elif abs(ficha - previa) > 1:
            if any(abs(ficha - t) <= 1 for t in tasas):
                log('Tarifas', f"alumno {a['ID_Alumno']}", 'Monto', ficha, previa,
                    f"La ficha ({ficha}) coincide con una tarifa anterior del historial: se toma como "
                    f"desactualizada y la vigente es la última cobrada ({previa}).")
            else:
                tarifas.append({'ID_Alumno': a['ID_Alumno'], 'Desde': corte.isoformat(), 'Tipo': 'Hora',
                                'Monto': ficha, 'Nota': 'inferida (ficha actual, distinta del historial)'})
    for i, t in enumerate(tarifas, start=1):
        t['ID_Tarifa'] = i
    log('Tarifas', '', '(hoja nueva)', '', f"{len(tarifas)} filas",
        "Una fila por cada cambio de $/hora efectivo (MontoCobrado / HorasReales) por alumno, con Desde = "
        "primera clase a ese precio. Se excluyen las clases con Revisar = Sí.")

    # ── Alumnos: Estado y PrecioHora ──
    limite = pd.Timestamp(corte - timedelta(days=DIAS_INACTIVIDAD))
    for a in alumnos:
        ultimas = df[df['ID_Alumno'] == a['ID_Alumno']]['_fecha']
        ultima = ultimas.max() if not ultimas.empty else None
        if a['ID_Alumno'] in ACTIVOS_FORZADOS:
            nuevo, motivo = 'Activo', ACTIVOS_FORZADOS[a['ID_Alumno']]
        elif ultima is not None and ultima >= limite:
            nuevo, motivo = 'Activo', f"Última clase {ultima.date().isoformat()}"
        else:
            desde = ultima.date().isoformat() if ultima is not None else 'nunca'
            nuevo, motivo = 'Pausado', f"Sin clases en los últimos {DIAS_INACTIVIDAD} días (última: {desde})"
        if a['Estado'] != nuevo:
            log('Alumnos', a['ID_Alumno'], 'Estado', a['Estado'], nuevo, motivo)
            a['Estado'] = nuevo
        vig = [t for t in tarifas if t['ID_Alumno'] == a['ID_Alumno'] and t['Desde'] <= corte.isoformat()]
        if vig:
            monto = sorted(vig, key=lambda t: (t['Desde'], t['ID_Tarifa']))[-1]['Monto']
            if abs(monto - a['PrecioHora']) > 0.5:
                log('Alumnos', a['ID_Alumno'], 'PrecioHora', a['PrecioHora'], monto,
                    'Espejo de la tarifa vigente al ' + corte.strftime('%d/%m/%Y'))
                a['PrecioHora'] = monto

    # ── Totales por mes (antes = Total original, después = MontoCobrado) ──
    originales = {parse_id(r['ID_Clase']): num(r['Total']) for r in clases_raw}
    df['TotalOriginal'] = df['ID_Clase'].map(originales)
    df['Mes'] = df['_fecha'].dt.strftime('%Y-%m')
    totales = df.groupby('Mes', as_index=False).agg(Clases=('ID_Clase', 'count'),
                                                    Antes=('TotalOriginal', 'sum'),
                                                    Despues=('MontoCobrado', 'sum'))

    # ── Escritura ──
    out = openpyxl.Workbook()
    out.remove(out.active)

    def hoja(nombre, columnas, filas, texto_cols=()):
        ws = out.create_sheet(nombre)
        ws.append(columnas)
        for f in filas:
            ws.append([f.get(c, '') if f.get(c) is not None else '' for c in columnas])
        # Las celdas de texto quedan con formato texto para que Sheets no las convierta
        for col_idx, col in enumerate(columnas, start=1):
            if col in texto_cols:
                for (celda,) in ws.iter_rows(min_row=2, min_col=col_idx, max_col=col_idx):
                    celda.number_format = '@'
                    if celda.value is not None and not isinstance(celda.value, str):
                        celda.value = str(celda.value)
        return ws

    hoja('Clases', COLUMNAS_CLASES, clases,
         texto_cols={'Fecha', 'Alumno', 'Materia', '¿Pagada?', 'Método', 'Revisar', 'NotaRevision'})
    hoja('Alumnos', COLUMNAS_ALUMNOS, alumnos, texto_cols=COLUMNAS_TEXTO_ALUMNOS)
    hoja('Tarifas', COLUMNAS_TARIFAS, tarifas, texto_cols={'Desde', 'Tipo', 'Nota'})
    hoja('IPC', COLUMNAS_IPC, [], texto_cols={'Mes'})

    ws = out.create_sheet('LEEME')
    ws.append([f"Corrección de datos generada por scripts/corregir_datos.py · fecha de corte "
               f"{corte.strftime('%d/%m/%Y')} · fuente: {Path(entrada).name}"])
    ws.append([])
    ws.append(['Hoja', 'ID', 'Columna', 'Antes', 'Después', 'Motivo'])
    for fila in leeme:
        ws.append(fila)
    ws.append([])
    ws.append(['Totales por mes'])
    ws.append(['Mes', 'Clases', 'Total antes', 'MontoCobrado después', 'Diferencia'])
    for _, t in totales.iterrows():
        ws.append([t['Mes'], int(t['Clases']), float(t['Antes']), float(t['Despues']),
                   float(t['Despues'] - t['Antes'])])
    ws.append(['TOTAL', int(totales['Clases'].sum()), float(totales['Antes'].sum()),
               float(totales['Despues'].sum()), float(totales['Despues'].sum() - totales['Antes'].sum())])
    out.save(salida)
    return leeme, totales, tarifas, alumnos


def main():
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument('--entrada', default=str(RAIZ / 'Clases Particulares.xlsx'))
    p.add_argument('--salida', default=str(RAIZ / 'Clases Particulares - corregido.xlsx'))
    p.add_argument('--corte', default='2026-10-26', help='Fecha de corte para Estado y tarifa vigente')
    args = p.parse_args()
    if Path(args.entrada).resolve() == Path(args.salida).resolve():
        sys.exit("La salida no puede ser el archivo original.")
    leeme, totales, tarifas, alumnos = corregir(args.entrada, args.salida, date.fromisoformat(args.corte))
    print(f"OK → {args.salida}")
    print(f"{len(leeme)} entradas en LEEME · {len(tarifas)} tarifas · {len(alumnos)} alumnos")
    print(totales.to_string(index=False))


if __name__ == '__main__':
    main()
